"""Resumable BRACS WSI inventory and pathology feature-bag extraction."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import time
import urllib.request
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import torch
from openpyxl import load_workbook
from PIL import Image
from torch.utils.data import DataLoader, Dataset
from torchvision.transforms import v2

from pathlab_ai_data.bracs import (
    OFFICIAL_CLASS_COUNTS,
    OFFICIAL_SPLIT_CLASS_COUNTS,
    OFFICIAL_SPLIT_COUNTS,
)
from pathlab_ai_data.bracs_roi import (
    OFFICIAL_ROI_PATIENT_COUNTS,
    OFFICIAL_ROI_SLIDE_COUNTS,
)

from .core import COARSE_GROUP, LABELS, sha256

BRACS_FTP_ROOT = "ftp://histoimage.na.icar.cnr.it/BRACS_WSI"
ENCODER_ID = "kaiko-vits16"
ENCODER_RELEASE = "0.0.1"
ENCODER_REPOSITORY = "kaiko-ai/towards_large_pathology_fms:main"
ENCODER_SOURCE = f"{ENCODER_REPOSITORY}:vits16"
ENCODER_LICENSE = "Kaiko non-commercial research license"
ENCODER_HUBCONF_SHA256 = (
    "b0f5dd8600126c505870e14c548fc94ee6963717e6ca080453ee12bd64a09e6a"
)
ENCODER_WEIGHTS_SHA256 = (
    "4a117a8138420036ef2319122c8cdc5f7c43ef3f621960a62770c8a920accb0e"
)
_FTP_LINE = re.compile(
    r"^\S+\s+\d+\s+\S+\s+\S+\s+(?P<size>\d+)\s+"
    r"\S+\s+\d+\s+\S+\s+(?P<name>.+)$"
)


@dataclass(frozen=True)
class BagExtractionConfig:
    inventory_path: Path
    bags_root: Path
    temporary_root: Path
    splits: tuple[str, ...] = ("train", "validation", "test")
    max_tiles: int = 128
    target_mpp: float = 0.5
    batch_size: int = 16
    threads: int = 6
    limit: int | None = None
    slide_ids: tuple[str, ...] = ()
    keep_downloads: bool = False
    rebuild_manifest: bool = True


class _ImageDataset(Dataset[torch.Tensor]):
    def __init__(self, images: list[Image.Image]) -> None:
        self.images = images
        self.transform = v2.Compose(
            [
                v2.ToImage(),
                v2.Resize(224, antialias=True),
                v2.CenterCrop(224),
                v2.ToDtype(torch.float32, scale=True),
                v2.Normalize(mean=(0.5, 0.5, 0.5), std=(0.5, 0.5, 0.5)),
            ]
        )

    def __len__(self) -> int:
        return len(self.images)

    def __getitem__(self, index: int) -> torch.Tensor:
        return self.transform(self.images[index])


def parse_ftp_listing(text: str) -> dict[str, int]:
    """Parse the Unix LIST form exposed by the official BRACS FTP server."""

    result: dict[str, int] = {}
    for line in text.splitlines():
        match = _FTP_LINE.match(line.strip())
        if not match:
            continue
        name = match.group("name").strip()
        if name.lower().endswith(".svs"):
            result[name] = int(match.group("size"))
    return result


def build_bracs_wsi_inventory(summary_path: Path, output_path: Path) -> dict[str, Any]:
    """Join official spreadsheet metadata to remote FTP identities and sizes."""

    metadata = _read_wsi_summary(summary_path.resolve())
    remote: dict[str, tuple[str, int]] = {}
    for split in ("train", "validation", "test"):
        remote_split = "val" if split == "validation" else split
        for label in LABELS:
            group = f"Group_{COARSE_GROUP[label]}"
            folder = f"Type_{label}"
            url = f"{BRACS_FTP_ROOT}/{remote_split}/{group}/{folder}/"
            with urllib.request.urlopen(url, timeout=60) as response:
                listing = response.read().decode("utf-8", errors="replace")
            for filename, size in parse_ftp_listing(listing).items():
                slide_id = filename.removesuffix(".svs")
                if slide_id in remote:
                    raise ValueError(f"duplicate remote WSI: {slide_id}")
                remote[slide_id] = (url + filename, size)
    missing = sorted(set(metadata) - set(remote))
    extra = sorted(set(remote) - set(metadata))
    if missing or extra:
        raise ValueError(
            f"BRACS inventory mismatch: missing={missing[:5]} extra={extra[:5]}"
        )
    rows = []
    for slide_id, values in metadata.items():
        url, size = remote[slide_id]
        rows.append(
            {
                "slide_id": slide_id,
                "patient_id": str(values["patient_id"]),
                "label": values["label"],
                "coarse_group": COARSE_GROUP[values["label"]],
                "split": values["split"],
                "source_split": values["source_split"],
                "roi_count": int(values["roi_count"]),
                "source_url": url,
                "source_size_bytes": size,
            }
        )
    rows.sort(
        key=lambda row: (_split_order(row["split"]), _slide_number(row["slide_id"]))
    )
    _validate_inventory(rows)
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    temporary.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in rows),
        encoding="utf-8",
    )
    temporary.replace(output_path)
    report = {
        "schema_version": 1,
        "slides": len(rows),
        "patients": len({row["patient_id"] for row in rows}),
        "total_bytes": sum(row["source_size_bytes"] for row in rows),
        "split_counts": dict(Counter(row["split"] for row in rows)),
        "class_counts": dict(Counter(row["label"] for row in rows)),
        "inventory_sha256": sha256(output_path),
        "source": BRACS_FTP_ROOT,
    }
    output_path.with_suffix(".summary.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def select_bracs_wsi_cohort(
    inventory_path: Path,
    output_path: Path,
    *,
    train_per_class: int,
    validation_per_class: int,
    test_per_class: int = 0,
) -> dict[str, Any]:
    """Create an explicit size-bounded balanced development cohort."""

    if train_per_class < 1 or validation_per_class < 1 or test_per_class < 0:
        raise ValueError("cohort counts must be positive, except test may be zero")
    rows = [
        json.loads(line)
        for line in inventory_path.resolve().read_text(encoding="utf-8").splitlines()
        if line
    ]
    requested = {
        "train": train_per_class,
        "validation": validation_per_class,
        "test": test_per_class,
    }
    selected: list[dict[str, Any]] = []
    for split, count in requested.items():
        if count == 0:
            continue
        for label in LABELS:
            candidates = sorted(
                [
                    row
                    for row in rows
                    if row["split"] == split and row["label"] == label
                ],
                key=lambda row: (
                    row["source_size_bytes"],
                    _slide_number(row["slide_id"]),
                ),
            )
            if len(candidates) < count:
                raise ValueError(f"not enough {split}/{label} slides for cohort")
            selected.extend(candidates[:count])
    selected.sort(
        key=lambda row: (_split_order(row["split"]), _slide_number(row["slide_id"]))
    )
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in selected),
        encoding="utf-8",
    )
    report = {
        "schema_version": 1,
        "selection": "smallest source files within every class and split",
        "selection_bias": (
            "Resource-bounded development cohort; results do not estimate the full "
            "BRACS WSI distribution."
        ),
        "requested_per_class": requested,
        "slides": len(selected),
        "patients": len({row["patient_id"] for row in selected}),
        "total_bytes": sum(row["source_size_bytes"] for row in selected),
        "split_counts": dict(Counter(row["split"] for row in selected)),
        "class_counts": dict(Counter(row["label"] for row in selected)),
        "source_inventory_sha256": sha256(inventory_path.resolve()),
        "cohort_sha256": sha256(output_path),
    }
    output_path.with_suffix(".summary.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def build_bracs_final_inventory(
    development_inventory: Path, full_inventory: Path, output_path: Path
) -> dict[str, Any]:
    """Combine frozen development rows with every official locked-test row."""

    development_inventory = development_inventory.resolve()
    full_inventory = full_inventory.resolve()
    development = [
        json.loads(line)
        for line in development_inventory.read_text(encoding="utf-8").splitlines()
        if line
    ]
    full = [
        json.loads(line)
        for line in full_inventory.read_text(encoding="utf-8").splitlines()
        if line
    ]
    if any(row["split"] == "test" for row in development):
        raise ValueError("development inventory must not contain test rows")
    by_id = {row["slide_id"]: row for row in full}
    for row in development:
        if row["slide_id"] not in by_id or row != by_id[row["slide_id"]]:
            raise ValueError(
                f"development row disagrees with full inventory: {row['slide_id']}"
            )
    locked_test = [row for row in full if row["split"] == "test"]
    if len(locked_test) != OFFICIAL_SPLIT_COUNTS["test"]:
        raise ValueError("full inventory does not contain the official locked test")
    combined = sorted(
        [*development, *locked_test],
        key=lambda row: (_split_order(row["split"]), _slide_number(row["slide_id"])),
    )
    patient_splits: dict[str, set[str]] = {}
    for row in combined:
        patient_splits.setdefault(str(row["patient_id"]), set()).add(row["split"])
    leaked = sorted(
        patient for patient, splits in patient_splits.items() if len(splits) > 1
    )
    if leaked:
        raise ValueError(f"patient leakage in final inventory: {leaked[:5]}")
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in combined),
        encoding="utf-8",
    )
    report = {
        "schema_version": 1,
        "development_selection": "frozen size-bounded balanced cohort",
        "test_selection": "all official corrected-split test slides",
        "slides": len(combined),
        "patients": len(patient_splits),
        "split_counts": dict(Counter(row["split"] for row in combined)),
        "test_class_counts": dict(Counter(row["label"] for row in locked_test)),
        "test_bytes": sum(int(row["source_size_bytes"]) for row in locked_test),
        "development_inventory_sha256": sha256(development_inventory),
        "full_inventory_sha256": sha256(full_inventory),
        "final_inventory_sha256": sha256(output_path),
    }
    output_path.with_suffix(".summary.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def build_bracs_development_inventory(
    full_inventory: Path, output_path: Path
) -> dict[str, Any]:
    """Freeze every patient-safe BRACS train/validation row without test rows."""

    full_inventory = full_inventory.resolve()
    rows = [
        json.loads(line)
        for line in full_inventory.read_text(encoding="utf-8").splitlines()
        if line
    ]
    development = [row for row in rows if row["split"] in {"train", "validation"}]
    split_counts = Counter(row["split"] for row in development)
    if dict(split_counts) != {"train": 392, "validation": 68}:
        raise ValueError(f"full development split mismatch: {dict(split_counts)}")
    if {row["label"] for row in development} != set(LABELS):
        raise ValueError("full development inventory does not cover every class")
    patient_splits: dict[str, set[str]] = {}
    for row in development:
        patient_splits.setdefault(str(row["patient_id"]), set()).add(row["split"])
    leaked = sorted(
        patient for patient, splits in patient_splits.items() if len(splits) > 1
    )
    if leaked:
        raise ValueError(f"patient leakage in development inventory: {leaked[:5]}")
    development.sort(
        key=lambda row: (_split_order(row["split"]), _slide_number(row["slide_id"]))
    )
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in development),
        encoding="utf-8",
    )
    report = {
        "schema_version": 1,
        "selection": "all patient-safe BRACS development slides",
        "slides": len(development),
        "patients": len(patient_splits),
        "split_counts": dict(split_counts),
        "class_counts": dict(Counter(row["label"] for row in development)),
        "source_bytes": sum(int(row["source_size_bytes"]) for row in development),
        "full_inventory_sha256": sha256(full_inventory),
        "development_inventory_sha256": sha256(output_path),
    }
    output_path.with_suffix(".summary.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def shard_bracs_wsi_inventory(
    inventory_path: Path, output_directory: Path, *, shards: int
) -> dict[str, Any]:
    """Create deterministic byte-balanced, disjoint extraction inventories."""

    if shards < 1:
        raise ValueError("shards must be at least one")
    inventory_path = inventory_path.resolve()
    rows = [
        json.loads(line)
        for line in inventory_path.read_text(encoding="utf-8").splitlines()
        if line
    ]
    if len({row["slide_id"] for row in rows}) != len(rows):
        raise ValueError("inventory contains duplicate slide IDs")
    partitions: list[list[dict[str, Any]]] = [[] for _ in range(shards)]
    totals = [0] * shards
    for row in sorted(
        rows, key=lambda item: (-int(item["source_size_bytes"]), item["slide_id"])
    ):
        index = min(range(shards), key=lambda item: (totals[item], item))
        partitions[index].append(row)
        totals[index] += int(row["source_size_bytes"])
    output_directory = output_directory.resolve()
    output_directory.mkdir(parents=True, exist_ok=True)
    records = []
    for index, partition in enumerate(partitions):
        partition.sort(
            key=lambda row: (
                _split_order(row["split"]),
                _slide_number(row["slide_id"]),
            )
        )
        path = output_directory / f"shard-{index}.jsonl"
        path.write_text(
            "".join(json.dumps(row, sort_keys=True) + "\n" for row in partition),
            encoding="utf-8",
        )
        records.append(
            {
                "shard": index,
                "slides": len(partition),
                "source_bytes": totals[index],
                "inventory": str(path),
                "inventory_sha256": sha256(path),
            }
        )
    report = {
        "schema_version": 1,
        "source_inventory_sha256": sha256(inventory_path),
        "shards": records,
        "slides": sum(len(partition) for partition in partitions),
        "source_bytes": sum(totals),
    }
    (output_directory / "shards.summary.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def _read_wsi_summary(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["WSI_Information"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        columns = {
            str(value).strip().lower(): index
            for index, value in enumerate(headers)
            if value is not None
        }
        required = {"wsi filename", "patient id", "roi", "wsi label", "set"}
        normalized = {name.replace(" ", ""): index for name, index in columns.items()}
        required_normalized = {name.replace(" ", "") for name in required}
        missing = required_normalized - normalized.keys()
        if missing:
            raise ValueError(f"WSI_Information missing columns: {sorted(missing)}")
        result: dict[str, dict[str, Any]] = {}
        split_names = {
            "training": "train",
            "train": "train",
            "validation": "validation",
            "val": "validation",
            "testing": "test",
            "test": "test",
        }
        for row_number, row in enumerate(rows, start=2):
            raw_id = row[normalized["wsifilename"]]
            if raw_id is None or not str(raw_id).strip():
                continue
            slide_id = str(raw_id).strip()
            label = str(row[normalized["wsilabel"]]).strip().upper()
            split_raw = str(row[normalized["set"]]).strip().lower()
            if label not in LABELS or split_raw not in split_names:
                raise ValueError(f"invalid WSI metadata at row {row_number}")
            if slide_id in result:
                raise ValueError(f"duplicate WSI metadata: {slide_id}")
            source_split = split_names[split_raw]
            result[slide_id] = {
                "patient_id": str(row[normalized["patientid"]]).strip(),
                "label": label,
                "split": source_split,
                "source_split": source_split,
                "roi_count": int(row[normalized["roi"]] or 0),
            }
        patient_splits: dict[str, set[str]] = {}
        for values in result.values():
            patient_splits.setdefault(values["patient_id"], set()).add(values["split"])
        for patient, splits in patient_splits.items():
            if len(splits) == 1:
                continue
            if "test" in splits:
                raise ValueError(f"patient {patient} crosses the official test split")
            if splits != {"train", "validation"}:
                raise ValueError(
                    f"unsupported split overlap for patient {patient}: {splits}"
                )
            for values in result.values():
                if values["patient_id"] == patient:
                    values["split"] = "validation"
        return result
    finally:
        workbook.close()


def extract_bracs_wsi_bags(config: BagExtractionConfig) -> dict[str, Any]:
    """Download each WSI transiently and persist only compact tile feature bags."""

    if config.max_tiles < 8:
        raise ValueError("max_tiles must be at least 8")
    inventory_path = config.inventory_path.resolve()
    inventory = [
        json.loads(line)
        for line in inventory_path.read_text(encoding="utf-8").splitlines()
        if line
    ]
    selected = [row for row in inventory if row["split"] in config.splits]
    if config.slide_ids:
        requested = set(config.slide_ids)
        selected = [row for row in selected if row["slide_id"] in requested]
        found = {row["slide_id"] for row in selected}
        if found != requested:
            raise ValueError(
                f"slide IDs absent from selected splits: {sorted(requested - found)}"
            )
    if config.limit is not None:
        selected = selected[: config.limit]
    bags_root = config.bags_root.resolve()
    temporary_root = config.temporary_root.resolve()
    bags_root.mkdir(parents=True, exist_ok=True)
    temporary_root.mkdir(parents=True, exist_ok=True)
    model = load_kaiko_encoder()
    torch.set_num_threads(max(1, config.threads))
    completed = 0
    skipped = 0
    failures: list[dict[str, str]] = []
    for row in selected:
        bag_path, metadata_path = _bag_paths(bags_root, row)
        if _valid_existing_bag(bag_path, metadata_path, row, config):
            skipped += 1
            print(
                json.dumps(
                    {
                        "event": "bag_skipped",
                        "slide_id": row["slide_id"],
                        "progress": completed + skipped + len(failures),
                        "total": len(selected),
                    }
                ),
                flush=True,
            )
            continue
        bag_path.parent.mkdir(parents=True, exist_ok=True)
        download = temporary_root / f"{row['slide_id']}.svs"
        bag_written = False
        download_completed = False
        try:
            _download(row["source_url"], download, int(row["source_size_bytes"]))
            download_completed = True
            result = _extract_one(download, row, model, config)
            _atomic_npz(
                bag_path,
                features=result.pop("features"),
                coordinates=result.pop("coordinates"),
                tissue_scores=result.pop("tissue_scores"),
            )
            metadata_path.write_text(
                json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8"
            )
            bag_written = True
            completed += 1
            print(
                json.dumps(
                    {
                        "event": "bag_completed",
                        "slide_id": row["slide_id"],
                        "progress": completed + skipped + len(failures),
                        "total": len(selected),
                    }
                ),
                flush=True,
            )
        except Exception as error:  # noqa: BLE001 - isolate and report per-slide failures
            failures.append({"slide_id": row["slide_id"], "error": str(error)})
            print(
                json.dumps(
                    {
                        "event": "bag_failed",
                        "slide_id": row["slide_id"],
                        "error": str(error),
                        "progress": completed + skipped + len(failures),
                        "total": len(selected),
                    }
                ),
                flush=True,
            )
        finally:
            if bag_written and download.exists() and not config.keep_downloads:
                download.unlink()
            elif download_completed and download.exists() and not config.keep_downloads:
                # A byte-complete file that failed decoding/extraction must not be
                # resumed forever at EOF. It is a pipeline-owned transient copy and
                # remains recoverable from the official source.
                download.unlink()
        if config.rebuild_manifest:
            _rebuild_bags_manifest(bags_root, inventory_path)
    manifest_path = (
        _rebuild_bags_manifest(bags_root, inventory_path)
        if config.rebuild_manifest
        else None
    )
    return {
        "schema_version": 1,
        "selected": len(selected),
        "completed": completed,
        "skipped": skipped,
        "failed": len(failures),
        "failures": failures,
        "bags_manifest": str(manifest_path) if manifest_path else None,
        "bags_manifest_sha256": sha256(manifest_path) if manifest_path else None,
    }


def build_bracs_roi_slide_bags(
    views_root: Path, roi_manifest: Path, feature_cache: Path, bags_root: Path
) -> dict[str, Any]:
    """Build patient-safe slide bags from public ROI features and WSI labels."""

    views_root = views_root.resolve()
    roi_manifest = roi_manifest.resolve()
    feature_cache = feature_cache.resolve()
    bags_root = bags_root.resolve()
    records = [
        json.loads(line)
        for line in (views_root / "views.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
        if line
    ]
    by_roi = {record["roi_id"]: record for record in records}
    source_records = {
        record["roi_id"]: record
        for record in (
            json.loads(line)
            for line in roi_manifest.read_text(encoding="utf-8").splitlines()
            if line
        )
    }
    if set(source_records) != set(by_roi):
        raise ValueError("clean ROI manifest and views manifest identities disagree")
    loaded = np.load(feature_cache, allow_pickle=False)
    if str(loaded["views_manifest_sha256"].item()) != sha256(
        views_root / "views.jsonl"
    ):
        raise ValueError("feature cache does not match views manifest")
    roi_features: dict[str, list[np.ndarray]] = {}
    for roi_id, features in zip(loaded["roi_id"].astype(str), loaded["features"]):
        if roi_id not in by_roi:
            raise ValueError(f"feature cache contains unknown ROI: {roi_id}")
        roi_features.setdefault(roi_id, []).append(features.astype(np.float32))
    slide_features: dict[str, list[np.ndarray]] = {}
    slide_metadata: dict[str, dict[str, Any]] = {}
    for roi_id, values in roi_features.items():
        record = by_roi[roi_id]
        source_record = source_records[roi_id]
        slide_id = record["slide_id"]
        slide_features.setdefault(slide_id, []).append(np.mean(values, axis=0))
        metadata = {
            "slide_id": slide_id,
            "patient_id": str(record["patient_id"]),
            "label": source_record["wsi_label"],
            "coarse_group": COARSE_GROUP[source_record["wsi_label"]],
            "split": "validation" if record["split"] == "val" else record["split"],
            "encoder": ENCODER_ID,
            "source": "official BRACS ROI features grouped by WSI",
        }
        if slide_id in slide_metadata and slide_metadata[slide_id] != metadata:
            raise ValueError(f"inconsistent WSI metadata for {slide_id}")
        slide_metadata[slide_id] = metadata
    bags_root.mkdir(parents=True, exist_ok=True)
    manifest: list[dict[str, Any]] = []
    for slide_id in sorted(slide_features, key=_slide_number):
        metadata = slide_metadata[slide_id]
        base = bags_root / ENCODER_ID / metadata["split"] / metadata["label"] / slide_id
        base.parent.mkdir(parents=True, exist_ok=True)
        bag_path = base.with_suffix(".npz")
        features = np.vstack(slide_features[slide_id]).astype(np.float16)
        _atomic_npz(
            bag_path,
            features=features,
            coordinates=np.zeros((len(features), 2), dtype=np.int32),
            tissue_scores=np.ones(len(features), dtype=np.float32),
        )
        row = {
            **metadata,
            "tile_count": len(features),
            "feature_dimension": int(features.shape[1]),
            "bag_path": bag_path.relative_to(bags_root).as_posix(),
            "bag_sha256": sha256(bag_path),
        }
        base.with_suffix(".json").write_text(
            json.dumps(row, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        manifest.append(row)
    split_counts = Counter(row["split"] for row in manifest)
    expected_slides = {
        "train": OFFICIAL_ROI_SLIDE_COUNTS["train"],
        "validation": OFFICIAL_ROI_SLIDE_COUNTS["val"],
        "test": OFFICIAL_ROI_SLIDE_COUNTS["test"],
    }
    if dict(split_counts) != expected_slides:
        raise ValueError(f"ROI-derived slide split mismatch: {dict(split_counts)}")
    patient_splits: dict[str, set[str]] = {}
    for row in manifest:
        patient_splits.setdefault(row["patient_id"], set()).add(row["split"])
    if any(len(splits) != 1 for splits in patient_splits.values()):
        raise ValueError("patient leakage in ROI-derived slide bags")
    patient_counts = Counter(next(iter(splits)) for splits in patient_splits.values())
    expected_patients = {
        "train": OFFICIAL_ROI_PATIENT_COUNTS["train"],
        "validation": OFFICIAL_ROI_PATIENT_COUNTS["val"],
        "test": OFFICIAL_ROI_PATIENT_COUNTS["test"],
    }
    if dict(patient_counts) != expected_patients:
        raise ValueError(f"ROI-derived patient split mismatch: {dict(patient_counts)}")
    manifest_path = bags_root / "bags.jsonl"
    manifest_path.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in manifest),
        encoding="utf-8",
    )
    report = {
        "schema_version": 1,
        "slides": len(manifest),
        "patients": len(patient_splits),
        "split_counts": dict(split_counts),
        "patient_counts": dict(patient_counts),
        "feature_cache_sha256": sha256(feature_cache),
        "roi_manifest_sha256": sha256(roi_manifest),
        "bags_manifest_sha256": sha256(manifest_path),
    }
    (bags_root / "bag_config.json").write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return report


def load_kaiko_encoder() -> torch.nn.Module:
    repository = Path(torch.hub.get_dir()) / "kaiko-ai_towards_large_pathology_fms_main"
    hubconf = repository / "hubconf.py"
    if not hubconf.is_file() or sha256(hubconf) != ENCODER_HUBCONF_SHA256:
        raise ValueError("Kaiko loader-code checksum mismatch")
    model = torch.hub.load(
        str(repository),
        "vits16",
        source="local",
    )
    checkpoint = Path(torch.hub.get_dir()) / "checkpoints" / "vits16.pth"
    if not checkpoint.is_file() or sha256(checkpoint) != ENCODER_WEIGHTS_SHA256:
        raise ValueError("Kaiko ViT-S/16 checkpoint checksum mismatch")
    model.eval()
    for parameter in model.parameters():
        parameter.requires_grad_(False)
    return model


def extract_kaiko_slide_features(
    slide_path: Path,
    *,
    target_mpp: float = 0.5,
    max_tiles: int = 128,
    batch_size: int = 16,
    model: torch.nn.Module | None = None,
) -> dict[str, Any]:
    """Create an annotation-free Kaiko feature bag from a local WSI."""

    import tiffslide

    encoder = model if model is not None else load_kaiko_encoder()
    with tiffslide.TiffSlide(slide_path) as slide:
        width, height = slide.dimensions
        mpp = _slide_mpp(slide.properties)
        source_pixels = max(224, round(224 * target_mpp / mpp))
        coordinates, tissue_scores = select_tissue_tiles(
            slide, source_pixels, max_tiles
        )
        images = [
            slide.read_region(
                (int(x), int(y)), 0, (source_pixels, source_pixels)
            ).convert("RGB")
            for x, y in coordinates
        ]
    loader = DataLoader(
        _ImageDataset(images), batch_size=batch_size, shuffle=False, num_workers=0
    )
    chunks: list[np.ndarray] = []
    with torch.inference_mode():
        for batch in loader:
            output = encoder(batch)
            if isinstance(output, (tuple, list)):
                output = output[0]
            chunks.append(output.detach().cpu().numpy().astype(np.float16))
    features = np.vstack(chunks)
    if features.shape[0] != len(coordinates) or not np.isfinite(features).all():
        raise ValueError("encoder produced invalid feature bag")
    return {
        "features": features,
        "coordinates": np.asarray(coordinates, dtype=np.int32),
        "tissue_scores": np.asarray(tissue_scores, dtype=np.float32),
        "slide_width": width,
        "slide_height": height,
        "source_mpp": mpp,
        "target_mpp": target_mpp,
        "source_tile_pixels": source_pixels,
    }


def verify_slide_feature_repeatability(
    slide_path: Path,
    output_path: Path,
    *,
    target_mpp: float = 0.5,
    max_tiles: int = 16,
    batch_size: int = 8,
) -> dict[str, Any]:
    """Run two real-slide passes and persist exact deterministic evidence."""

    started = time.perf_counter()
    slide_path = slide_path.resolve()
    model = load_kaiko_encoder()
    first = extract_kaiko_slide_features(
        slide_path,
        target_mpp=target_mpp,
        max_tiles=max_tiles,
        batch_size=batch_size,
        model=model,
    )
    second = extract_kaiko_slide_features(
        slide_path,
        target_mpp=target_mpp,
        max_tiles=max_tiles,
        batch_size=batch_size,
        model=model,
    )
    coordinates_equal = bool(
        np.array_equal(first["coordinates"], second["coordinates"])
    )
    tissue_error = float(
        np.max(np.abs(first["tissue_scores"] - second["tissue_scores"]))
    )
    feature_error = float(
        np.max(
            np.abs(
                first["features"].astype(np.float32)
                - second["features"].astype(np.float32)
            )
        )
    )
    feature_hash = hashlib.sha256(first["features"].tobytes()).hexdigest()
    passed = coordinates_equal and tissue_error == 0.0 and feature_error == 0.0
    report = {
        "schema_version": 1,
        "status": "verified" if passed else "failed",
        "source": str(slide_path),
        "source_sha256": sha256(slide_path),
        "slide_width": first["slide_width"],
        "slide_height": first["slide_height"],
        "source_mpp": first["source_mpp"],
        "target_mpp": target_mpp,
        "tiles": len(first["features"]),
        "feature_dimension": int(first["features"].shape[1]),
        "feature_sha256": feature_hash,
        "coordinates_equal": coordinates_equal,
        "tissue_score_max_absolute_error": tissue_error,
        "feature_max_absolute_error": feature_error,
        "encoder_release": ENCODER_RELEASE,
        "encoder_hubconf_sha256": ENCODER_HUBCONF_SHA256,
        "encoder_weights_sha256": ENCODER_WEIGHTS_SHA256,
        "runtime_seconds": time.perf_counter() - started,
    }
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    if not passed:
        raise RuntimeError(f"real-slide feature repeatability failed: {report}")
    return report


def _download(url: str, destination: Path, expected_size: int) -> None:
    curl = shutil.which("curl") or shutil.which("curl.exe")
    if curl is None:
        raise RuntimeError("curl is required for resumable WSI downloads")
    if destination.exists() and destination.stat().st_size > expected_size:
        # A killed duplicate process can append beyond the remote size. Truncate only
        # this exact pipeline-owned temporary file, then let curl restart it cleanly.
        destination.write_bytes(b"")
    result = subprocess.run(
        [
            curl,
            "--fail",
            "--silent",
            "--show-error",
            "--retry",
            "5",
            "--retry-delay",
            "2",
            "--continue-at",
            "-",
            "--output",
            str(destination),
            url,
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        raise RuntimeError(
            f"download failed ({result.returncode}): {result.stderr.strip()}"
        )
    actual = destination.stat().st_size
    if actual != expected_size:
        raise ValueError(
            f"download size mismatch: expected={expected_size} actual={actual}"
        )


def _extract_one(
    slide_path: Path,
    row: dict[str, Any],
    model: torch.nn.Module,
    config: BagExtractionConfig,
) -> dict[str, Any]:
    source_hash = _file_hash(slide_path)
    extracted = extract_kaiko_slide_features(
        slide_path,
        target_mpp=config.target_mpp,
        max_tiles=config.max_tiles,
        batch_size=config.batch_size,
        model=model,
    )
    features = extracted.pop("features")
    coordinates = extracted.pop("coordinates")
    tissue_scores = extracted.pop("tissue_scores")
    return {
        "schema_version": 1,
        "slide_id": row["slide_id"],
        "patient_id": row["patient_id"],
        "label": row["label"],
        "coarse_group": row["coarse_group"],
        "split": row["split"],
        "source_url": row["source_url"],
        "source_size_bytes": row["source_size_bytes"],
        "source_sha256": source_hash,
        **extracted,
        "tile_count": len(coordinates),
        "encoder": ENCODER_ID,
        "encoder_source": ENCODER_SOURCE,
        "encoder_release": ENCODER_RELEASE,
        "encoder_hubconf_sha256": ENCODER_HUBCONF_SHA256,
        "encoder_weights_sha256": ENCODER_WEIGHTS_SHA256,
        "encoder_license": ENCODER_LICENSE,
        "feature_dimension": int(features.shape[1]),
        "features": features,
        "coordinates": np.asarray(coordinates, dtype=np.int32),
        "tissue_scores": np.asarray(tissue_scores, dtype=np.float32),
    }


def select_tissue_tiles(
    slide: Any, source_pixels: int, max_tiles: int
) -> tuple[list[tuple[int, int]], list[float]]:
    """Select deterministic, spatially distributed tissue tiles without annotations."""

    width, height = slide.dimensions
    thumbnail = slide.get_thumbnail((2048, 2048)).convert("RGB")
    pixels = np.asarray(thumbnail, dtype=np.uint8)
    channel_range = pixels.max(axis=2).astype(np.int16) - pixels.min(axis=2).astype(
        np.int16
    )
    brightness = pixels.mean(axis=2)
    mask = (channel_range >= 18) & (brightness >= 35) & (brightness <= 235)
    candidates: list[tuple[int, int, float]] = []
    stride = source_pixels
    thumb_w, thumb_h = thumbnail.size
    for y in range(0, max(1, height - source_pixels + 1), stride):
        for x in range(0, max(1, width - source_pixels + 1), stride):
            tx0 = min(thumb_w - 1, int(x / width * thumb_w))
            ty0 = min(thumb_h - 1, int(y / height * thumb_h))
            tx1 = max(tx0 + 1, min(thumb_w, int((x + source_pixels) / width * thumb_w)))
            ty1 = max(
                ty0 + 1, min(thumb_h, int((y + source_pixels) / height * thumb_h))
            )
            fraction = float(mask[ty0:ty1, tx0:tx1].mean())
            if fraction >= 0.20:
                candidates.append((x, y, fraction))
    if len(candidates) < 8:
        raise ValueError(f"too few tissue tiles: {len(candidates)}")
    if len(candidates) > max_tiles:
        indexes = np.linspace(0, len(candidates) - 1, max_tiles, dtype=int)
        candidates = [candidates[int(index)] for index in indexes]
    return (
        [(x, y) for x, y, _ in candidates],
        [score for _, _, score in candidates],
    )


def _slide_mpp(properties: dict[str, Any]) -> float:
    for key in ("tiffslide.mpp-x", "openslide.mpp-x"):
        try:
            value = float(properties[key])
            if 0.05 <= value <= 5.0:
                return value
        except (KeyError, TypeError, ValueError):
            pass
    return 0.25


def _bag_paths(root: Path, row: dict[str, Any]) -> tuple[Path, Path]:
    base = root / ENCODER_ID / row["split"] / row["label"] / row["slide_id"]
    return base.with_suffix(".npz"), base.with_suffix(".json")


def _valid_existing_bag(
    bag_path: Path,
    metadata_path: Path,
    row: dict[str, Any],
    config: BagExtractionConfig,
) -> bool:
    if not bag_path.is_file() or not metadata_path.is_file():
        return False
    try:
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        with np.load(bag_path, allow_pickle=False) as loaded:
            features = loaded["features"]
            coordinates = loaded["coordinates"]
            tissue_scores = loaded["tissue_scores"]
        width = int(metadata["slide_width"])
        height = int(metadata["slide_height"])
        source_pixels = int(metadata["source_tile_pixels"])
        return (
            metadata["slide_id"] == row["slide_id"]
            and metadata["source_size_bytes"] == row["source_size_bytes"]
            and metadata["target_mpp"] == config.target_mpp
            and metadata["tile_count"] <= config.max_tiles
            and features.shape
            == (metadata["tile_count"], metadata["feature_dimension"])
            and coordinates.shape == (metadata["tile_count"], 2)
            and tissue_scores.shape == (metadata["tile_count"],)
            and np.isfinite(features).all()
            and np.isfinite(coordinates).all()
            and np.isfinite(tissue_scores).all()
            and ((tissue_scores >= 0) & (tissue_scores <= 1)).all()
            and len(np.unique(coordinates, axis=0)) == metadata["tile_count"]
            and (coordinates >= 0).all()
            and (coordinates[:, 0] + source_pixels <= width).all()
            and (coordinates[:, 1] + source_pixels <= height).all()
        )
    except (OSError, KeyError, ValueError, json.JSONDecodeError):
        return False


def _atomic_npz(path: Path, **arrays: np.ndarray) -> None:
    with tempfile.NamedTemporaryFile(
        dir=path.parent, suffix=".npz", delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        np.savez_compressed(temporary, **arrays)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _rebuild_bags_manifest(root: Path, inventory_path: Path) -> Path:
    inventory = {
        row["slide_id"]: row
        for row in (
            json.loads(line)
            for line in inventory_path.read_text(encoding="utf-8").splitlines()
            if line
        )
    }
    records: list[dict[str, Any]] = []
    for metadata_path in sorted((root / ENCODER_ID).glob("*/*/*.json")):
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        slide_id = metadata["slide_id"]
        if slide_id not in inventory:
            raise ValueError(f"bag is absent from selected inventory: {slide_id}")
        source = inventory[slide_id]
        for key in ("patient_id", "label", "split", "source_size_bytes"):
            if str(metadata[key]) != str(source[key]):
                raise ValueError(
                    f"bag metadata disagrees with inventory for {slide_id}: {key}"
                )
        bag_path = metadata_path.with_suffix(".npz")
        if not bag_path.is_file():
            continue
        metadata["bag_path"] = bag_path.relative_to(root).as_posix()
        metadata["encoder_source"] = ENCODER_SOURCE
        metadata["encoder_release"] = ENCODER_RELEASE
        metadata["encoder_hubconf_sha256"] = ENCODER_HUBCONF_SHA256
        metadata["encoder_weights_sha256"] = ENCODER_WEIGHTS_SHA256
        metadata["encoder_license"] = ENCODER_LICENSE
        metadata["bag_sha256"] = sha256(bag_path)
        metadata["metadata_path"] = metadata_path.relative_to(root).as_posix()
        records.append(metadata)
    records.sort(
        key=lambda row: (_split_order(row["split"]), _slide_number(row["slide_id"]))
    )
    path = root / "bags.jsonl"
    temporary = path.with_suffix(".jsonl.tmp")
    temporary.write_text(
        "".join(json.dumps(row, sort_keys=True) + "\n" for row in records),
        encoding="utf-8",
    )
    temporary.replace(path)
    (root / "bag_config.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "encoder": ENCODER_ID,
                "encoder_source": ENCODER_SOURCE,
                "encoder_release": ENCODER_RELEASE,
                "encoder_hubconf_sha256": ENCODER_HUBCONF_SHA256,
                "encoder_weights_sha256": ENCODER_WEIGHTS_SHA256,
                "encoder_license": ENCODER_LICENSE,
                "inventory_sha256": sha256(inventory_path),
                "bags": len(records),
            },
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    return path


def verify_bracs_wsi_bags(inventory_path: Path, bags_root: Path) -> dict[str, Any]:
    """Fail closed unless the feature cohort exactly matches its inventory."""

    inventory_path = inventory_path.resolve()
    bags_root = bags_root.resolve()
    inventory = {
        row["slide_id"]: row
        for row in (
            json.loads(line)
            for line in inventory_path.read_text(encoding="utf-8").splitlines()
            if line
        )
    }
    manifest_path = bags_root / "bags.jsonl"
    records = [
        json.loads(line)
        for line in manifest_path.read_text(encoding="utf-8").splitlines()
        if line
    ]
    identities = [record["slide_id"] for record in records]
    duplicates = sorted(
        slide_id for slide_id, count in Counter(identities).items() if count > 1
    )
    if duplicates:
        raise ValueError(f"duplicate bags: {duplicates[:5]}")
    actual = set(identities)
    expected = set(inventory)
    if actual != expected:
        raise ValueError(
            "bag cohort identity mismatch: "
            f"missing={sorted(expected - actual)[:5]} extra={sorted(actual - expected)[:5]}"
        )
    patient_splits: dict[str, set[str]] = {}
    split_counts: Counter[str] = Counter()
    class_counts: Counter[str] = Counter()
    tile_counts: list[int] = []
    dimensions: set[int] = set()
    source_bytes = 0
    bag_bytes = 0
    for record in records:
        source = inventory[record["slide_id"]]
        if (
            record["encoder"] != ENCODER_ID
            or record["encoder_release"] != ENCODER_RELEASE
            or record["encoder_hubconf_sha256"] != ENCODER_HUBCONF_SHA256
            or record["encoder_weights_sha256"] != ENCODER_WEIGHTS_SHA256
        ):
            raise ValueError(f"encoder provenance mismatch: {record['slide_id']}")
        for key in ("patient_id", "label", "split", "source_size_bytes"):
            if str(record[key]) != str(source[key]):
                raise ValueError(
                    f"bag metadata mismatch for {record['slide_id']}: {key}"
                )
        bag_path = bags_root / record["bag_path"]
        if sha256(bag_path) != record["bag_sha256"]:
            raise ValueError(f"bag checksum mismatch: {record['slide_id']}")
        with np.load(bag_path, allow_pickle=False) as loaded:
            features = loaded["features"]
            coordinates = loaded["coordinates"]
            tissue_scores = loaded["tissue_scores"]
        source_bytes += int(record["source_size_bytes"])
        bag_bytes += bag_path.stat().st_size
        tile_count = int(record["tile_count"])
        dimension = int(record["feature_dimension"])
        width = int(record["slide_width"])
        height = int(record["slide_height"])
        source_pixels = int(record["source_tile_pixels"])
        if not 8 <= tile_count <= 128:
            raise ValueError(f"tile count violates protocol: {record['slide_id']}")
        if float(record["target_mpp"]) != 0.5:
            raise ValueError(f"target MPP violates protocol: {record['slide_id']}")
        if width <= 0 or height <= 0 or source_pixels < 224:
            raise ValueError(f"invalid slide geometry: {record['slide_id']}")
        if features.shape != (tile_count, dimension):
            raise ValueError(f"feature shape mismatch: {record['slide_id']}")
        if coordinates.shape != (tile_count, 2):
            raise ValueError(f"coordinate shape mismatch: {record['slide_id']}")
        if tissue_scores.shape != (tile_count,):
            raise ValueError(f"tissue-score shape mismatch: {record['slide_id']}")
        if not all(
            np.isfinite(values).all()
            for values in (features, coordinates, tissue_scores)
        ):
            raise ValueError(f"non-finite bag values: {record['slide_id']}")
        if not ((tissue_scores >= 0) & (tissue_scores <= 1)).all():
            raise ValueError(f"invalid tissue scores: {record['slide_id']}")
        if (tissue_scores < 0.20).any():
            raise ValueError(f"tissue score violates protocol: {record['slide_id']}")
        if len(np.unique(coordinates, axis=0)) != tile_count:
            raise ValueError(f"duplicate tile coordinates: {record['slide_id']}")
        if (
            (coordinates < 0).any()
            or (coordinates[:, 0] + source_pixels > width).any()
            or (coordinates[:, 1] + source_pixels > height).any()
        ):
            raise ValueError(f"tile coordinates outside slide: {record['slide_id']}")
        patient_splits.setdefault(str(record["patient_id"]), set()).add(record["split"])
        split_counts[record["split"]] += 1
        class_counts[f"{record['split']}:{record['label']}"] += 1
        tile_counts.append(tile_count)
        dimensions.add(dimension)
    leaked = sorted(
        patient for patient, splits in patient_splits.items() if len(splits) > 1
    )
    if leaked:
        raise ValueError(f"patient leakage across bag splits: {leaked[:5]}")
    if dimensions != {384}:
        raise ValueError(f"unexpected Kaiko feature dimensions: {sorted(dimensions)}")
    report = {
        "schema_version": 1,
        "status": "verified",
        "bags": len(records),
        "patients": len(patient_splits),
        "split_counts": dict(sorted(split_counts.items())),
        "split_class_counts": dict(sorted(class_counts.items())),
        "feature_dimension": 384,
        "minimum_tiles": min(tile_counts),
        "maximum_tiles": max(tile_counts),
        "total_tile_vectors": sum(tile_counts),
        "mean_tiles_per_slide": float(np.mean(tile_counts)),
        "source_bytes": source_bytes,
        "feature_bag_bytes": bag_bytes,
        "source_to_feature_storage_ratio": source_bytes / bag_bytes,
        "inventory_sha256": sha256(inventory_path),
        "bags_manifest_sha256": sha256(manifest_path),
    }
    verification_path = bags_root / "verification.json"
    verification_path.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    report["verification_path"] = str(verification_path)
    return report


def _validate_inventory(rows: list[dict[str, Any]]) -> None:
    if len(rows) != sum(OFFICIAL_SPLIT_COUNTS.values()):
        raise ValueError(f"expected 547 slides, found {len(rows)}")
    split_counts = Counter(row["source_split"] for row in rows)
    if dict(split_counts) != dict(OFFICIAL_SPLIT_COUNTS):
        raise ValueError(f"official split mismatch: {dict(split_counts)}")
    class_counts = Counter(row["label"] for row in rows)
    if dict(class_counts) != dict(OFFICIAL_CLASS_COUNTS):
        raise ValueError(f"official class mismatch: {dict(class_counts)}")
    for split, expected in OFFICIAL_SPLIT_CLASS_COUNTS.items():
        actual = Counter(row["label"] for row in rows if row["source_split"] == split)
        if dict(actual) != dict(expected):
            raise ValueError(f"official {split} class mismatch: {dict(actual)}")
    source_patient_splits: dict[str, set[str]] = {}
    patient_splits: dict[str, set[str]] = {}
    for row in rows:
        source_patient_splits.setdefault(row["patient_id"], set()).add(
            row["source_split"]
        )
        patient_splits.setdefault(row["patient_id"], set()).add(row["split"])
    leaked = [patient for patient, splits in patient_splits.items() if len(splits) > 1]
    if leaked:
        raise ValueError(f"patient leakage in official inventory: {leaked[:5]}")
    source_patient_counts = Counter()
    for splits in source_patient_splits.values():
        for split in splits:
            source_patient_counts[split] += 1
    expected_source_patient_counts = {"train": 133, "validation": 26, "test": 31}
    if dict(source_patient_counts) != expected_source_patient_counts:
        raise ValueError(
            f"official patient split mismatch: {dict(source_patient_counts)}"
        )
    corrected_counts = Counter(row["split"] for row in rows)
    if dict(corrected_counts) != {"train": 392, "validation": 68, "test": 87}:
        raise ValueError(f"patient-safe split mismatch: {dict(corrected_counts)}")


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _slide_number(slide_id: str) -> int:
    return int(slide_id.split("_", 1)[1])


def _split_order(split: str) -> int:
    return {"train": 0, "validation": 1, "test": 2}[split]
