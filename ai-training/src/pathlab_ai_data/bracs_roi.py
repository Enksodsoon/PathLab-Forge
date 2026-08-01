"""Fail-closed preparation of the official BRACS ROI dataset."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError


LABELS = ("N", "PB", "UDH", "FEA", "ADH", "DCIS", "IC")
OFFICIAL_ROI_COUNT = 4_539
OFFICIAL_ROI_CLASS_COUNTS = {
    "N": 484,
    "PB": 836,
    "UDH": 517,
    "FEA": 756,
    "ADH": 507,
    "DCIS": 790,
    "IC": 649,
}
OFFICIAL_ROI_SPLIT_COUNTS = {"train": 3_657, "val": 312, "test": 570}
OFFICIAL_ROI_SPLIT_CLASS_COUNTS = {
    "train": {"N": 357, "PB": 714, "UDH": 389, "FEA": 624, "ADH": 387, "DCIS": 665, "IC": 521},
    "val": {"N": 46, "PB": 43, "UDH": 46, "FEA": 49, "ADH": 41, "DCIS": 40, "IC": 47},
    "test": {"N": 81, "PB": 79, "UDH": 82, "FEA": 83, "ADH": 79, "DCIS": 85, "IC": 81},
}
OFFICIAL_ROI_PATIENT_COUNTS = {"train": 106, "val": 15, "test": 30}
OFFICIAL_ROI_SLIDE_COUNTS = {"train": 281, "val": 37, "test": 69}

_ROI_NAME = re.compile(
    r"^(BRACS_(?P<slide_number>\d+))_(?P<label>N|PB|UDH|FEA|ADH|DCIS|IC)_(?P<roi_number>\d+)\.png$",
    re.IGNORECASE,
)
_LABEL_FOLDER = re.compile(r"^[0-6]_(N|PB|UDH|FEA|ADH|DCIS|IC)$", re.IGNORECASE)
_SPLIT_ALIASES = {
    "train": "train",
    "training": "train",
    "val": "val",
    "validation": "val",
    "test": "test",
    "testing": "test",
}


class RoiPreparationError(RuntimeError):
    """The ROI source cannot safely produce a training manifest."""


@dataclass(frozen=True)
class BracsRoiPreparationConfig:
    raw_root: Path
    summary_path: Path
    output_root: Path
    enforce_official_counts: bool = True
    minimum_dimension: int = 64
    minimum_file_bytes: int = 1_024


def prepare_bracs_roi(config: BracsRoiPreparationConfig) -> dict[str, Any]:
    """Decode, validate, hash, and index BRACS ROI images without copying pixels."""
    raw_root = config.raw_root.resolve()
    summary_path = config.summary_path.resolve()
    output_root = config.output_root.resolve()
    if not raw_root.is_dir():
        raise RoiPreparationError(f"ROI root does not exist: {raw_root}")
    if not summary_path.is_file():
        raise RoiPreparationError(f"BRACS summary does not exist: {summary_path}")
    if _is_within(output_root, raw_root):
        raise RoiPreparationError("output root must be outside the immutable raw dataset")

    metadata = _read_wsi_metadata(summary_path)
    images = sorted(path for path in raw_root.rglob("*.png") if path.is_file())
    if not images:
        raise RoiPreparationError("no ROI PNG images were found")
    if config.enforce_official_counts and len(images) != OFFICIAL_ROI_COUNT:
        raise RoiPreparationError(
            f"expected {OFFICIAL_ROI_COUNT} ROI images, found {len(images)}"
        )

    rows: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    content_hashes: dict[str, str] = {}
    for path in images:
        relative = path.relative_to(raw_root)
        split, folder_label = _path_facts(relative)
        match = _ROI_NAME.fullmatch(path.name)
        if not match:
            raise RoiPreparationError(f"invalid ROI filename: {relative.as_posix()}")
        label = match.group("label").upper()
        if label != folder_label:
            raise RoiPreparationError(f"filename/folder label mismatch: {relative.as_posix()}")
        slide_id = f"BRACS_{int(match.group('slide_number'))}"
        roi_number = int(match.group("roi_number"))
        roi_id = f"{slide_id}_{label}_{roi_number}"
        if roi_id in seen_ids:
            raise RoiPreparationError(f"duplicate ROI identity: {roi_id}")
        seen_ids.add(roi_id)
        slide = metadata.get(slide_id)
        if slide is None:
            raise RoiPreparationError(f"ROI has no WSI metadata row: {roi_id}")
        if slide["split"] != split:
            raise RoiPreparationError(
                f"ROI split disagrees with BRACS.xlsx for {roi_id}: {split} != {slide['split']}"
            )
        size = path.stat().st_size
        if size < config.minimum_file_bytes:
            raise RoiPreparationError(f"ROI file is unexpectedly small: {relative.as_posix()}")
        width, height, mode = _probe_png(path)
        if min(width, height) < config.minimum_dimension:
            raise RoiPreparationError(
                f"ROI dimensions are below {config.minimum_dimension}px: {relative.as_posix()} ({width}x{height})"
            )
        digest = _sha256(path)
        duplicate = content_hashes.get(digest)
        if duplicate is not None:
            raise RoiPreparationError(
                f"duplicate image content detected: {duplicate} and {relative.as_posix()}"
            )
        content_hashes[digest] = relative.as_posix()
        rows.append(
            {
                "roi_id": roi_id,
                "slide_id": slide_id,
                "patient_id": slide["patient_id"],
                "split": split,
                "label": label,
                "wsi_label": slide["wsi_label"],
                "roi_number": roi_number,
                "relative_path": relative.as_posix(),
                "file_size_bytes": size,
                "sha256": digest,
                "width": width,
                "height": height,
                "mode": mode,
            }
        )

    report = _validate_distribution(rows, config.enforce_official_counts)
    _write_output(output_root, rows, report, raw_root, summary_path)
    return report


def _read_wsi_metadata(summary_path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(summary_path, read_only=True, data_only=True)
    try:
        if "WSI_Information" not in workbook.sheetnames:
            raise RoiPreparationError("BRACS.xlsx is missing the WSI_Information sheet")
        sheet = workbook["WSI_Information"]
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            raise RoiPreparationError("WSI_Information is empty")
        names = {str(value).strip().lower(): index for index, value in enumerate(headers) if value}
        required = {"wsi filename", "patient id", "wsi label", "set"}
        missing = sorted(required - names.keys())
        if missing:
            raise RoiPreparationError(f"WSI_Information missing columns: {', '.join(missing)}")
        result: dict[str, dict[str, str]] = {}
        for row_number, values_row in enumerate(values, start=2):
            raw_id = values_row[names["wsi filename"]]
            if raw_id is None or str(raw_id).strip() == "":
                continue
            match = re.fullmatch(r"BRACS_(\d+)", str(raw_id).strip(), re.IGNORECASE)
            if not match:
                raise RoiPreparationError(f"invalid WSI identity at spreadsheet row {row_number}: {raw_id!r}")
            slide_id = f"BRACS_{int(match.group(1))}"
            split_raw = str(values_row[names["set"]]).strip().lower()
            split = _SPLIT_ALIASES.get(split_raw)
            if split is None:
                raise RoiPreparationError(f"invalid split at spreadsheet row {row_number}: {split_raw!r}")
            if slide_id in result:
                raise RoiPreparationError(f"duplicate WSI metadata row: {slide_id}")
            result[slide_id] = {
                "patient_id": str(values_row[names["patient id"]]).strip(),
                "wsi_label": str(values_row[names["wsi label"]]).strip().upper(),
                "split": split,
            }
        return result
    finally:
        workbook.close()


def _path_facts(relative: Path) -> tuple[str, str]:
    if len(relative.parts) != 3:
        raise RoiPreparationError(
            f"expected split/class/file ROI structure: {relative.as_posix()}"
        )
    split = _SPLIT_ALIASES.get(relative.parts[0].lower())
    label_match = _LABEL_FOLDER.fullmatch(relative.parts[1])
    if split is None or label_match is None:
        raise RoiPreparationError(f"invalid ROI folder structure: {relative.as_posix()}")
    return split, label_match.group(1).upper()


def _probe_png(path: Path) -> tuple[int, int, str]:
    try:
        with Image.open(path) as image:
            if image.format != "PNG":
                raise RoiPreparationError(f"file is not decoded as PNG: {path}")
            width, height, mode = image.width, image.height, image.mode
            image.verify()
            return width, height, mode
    except (OSError, ValueError, UnidentifiedImageError) as error:
        raise RoiPreparationError(f"cannot decode ROI PNG {path}: {error}") from error


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _validate_distribution(rows: list[dict[str, Any]], official: bool) -> dict[str, Any]:
    split_counts = Counter(row["split"] for row in rows)
    class_counts = Counter(row["label"] for row in rows)
    split_class = {
        split: dict(Counter(row["label"] for row in rows if row["split"] == split))
        for split in ("train", "val", "test")
    }
    patients: dict[str, set[str]] = defaultdict(set)
    slides: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        patients[row["split"]].add(row["patient_id"])
        slides[row["split"]].add(row["slide_id"])
    for left, right in (("train", "val"), ("train", "test"), ("val", "test")):
        overlap = patients[left] & patients[right]
        if overlap:
            raise RoiPreparationError(
                f"patient leakage between {left} and {right}: {sorted(overlap)[:5]}"
            )
    patient_counts = {key: len(patients[key]) for key in ("train", "val", "test")}
    slide_counts = {key: len(slides[key]) for key in ("train", "val", "test")}
    if official:
        checks = (
            (dict(class_counts), OFFICIAL_ROI_CLASS_COUNTS, "class"),
            (dict(split_counts), OFFICIAL_ROI_SPLIT_COUNTS, "split"),
            (split_class, OFFICIAL_ROI_SPLIT_CLASS_COUNTS, "split/class"),
            (patient_counts, OFFICIAL_ROI_PATIENT_COUNTS, "patient split"),
            (slide_counts, OFFICIAL_ROI_SLIDE_COUNTS, "slide split"),
        )
        for actual, expected, name in checks:
            if actual != expected:
                raise RoiPreparationError(f"{name} counts differ from official BRACS: {actual}")
    return {
        "dataset": "BRACS_RoI_latest_version",
        "images": len(rows),
        "bytes": sum(row["file_size_bytes"] for row in rows),
        "class_counts": dict(sorted(class_counts.items())),
        "split_counts": dict(sorted(split_counts.items())),
        "split_class_counts": split_class,
        "patient_counts": patient_counts,
        "slide_counts": slide_counts,
        "patient_disjoint": True,
        "duplicate_content_images": 0,
    }


def _write_output(
    output_root: Path,
    rows: list[dict[str, Any]],
    report: dict[str, Any],
    raw_root: Path,
    summary_path: Path,
) -> None:
    output_root.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{output_root.name}-", dir=output_root.parent))
    try:
        fields = list(rows[0].keys())
        with (staging / "manifest.csv").open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        with (staging / "manifest.jsonl").open("w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(row, sort_keys=True) + "\n")
        provenance = {
            **report,
            "source_url": "https://www.bracs.icar.cnr.it/download/",
            "source_dataset_path": "BRACS_RoI/latest_version",
            "license_declared_by_current_source": "CC-BY-NC-4.0",
            "raw_root": str(raw_root),
            "summary_path": str(summary_path),
            "raw_data_copied": False,
            "manifest_sha256": _sha256(staging / "manifest.csv"),
        }
        (staging / "provenance.json").write_text(
            json.dumps(provenance, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        if output_root.exists():
            shutil.rmtree(output_root)
        os.replace(staging, output_root)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


def _is_within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False
